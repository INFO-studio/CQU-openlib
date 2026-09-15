from __future__ import annotations

from collections import Counter
from contextlib import redirect_stderr, redirect_stdout
from io import BytesIO, StringIO
import hashlib
import json
from pathlib import Path
import random
import re
import unittest
from unittest.mock import patch
from urllib.parse import unquote
from zipfile import ZipFile

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

import clean_training_plans as plans


HEADERS = (
    "学院", "专业", "年级", "层次", "课程类别", "课程性质", "学科类别", "课程名称",
    "课程代码", "总学分", "总学时", "理论学时", "实验学时", "实践学时", "课外学时",
    "开课学期", "考核方式", "备注",
)


def row(**changes: plans.CellValue) -> plans.RawFields:
    fields: plans.RawFields = dict.fromkeys(HEADERS)
    fields.update({
        "学院": "测试学院", "专业": "测试专业", "年级": 2026, "层次": "本科",
        "课程类别": "公共基础课程", "课程性质": "必修", "课程名称": "测试课程",
        "课程代码": "TEST100", "总学分": 2, "开课学期": "1",
    })
    fields.update(changes)
    return fields


def sources(*rows: plans.RawFields) -> list[tuple[plans.RawFields, plans.Source]]:
    return [(fields, {"file": "a.xlsx", "sheet": "Sheet", "row": number}) for number, fields in enumerate(rows, start=3)]


def clean(*rows: plans.RawFields, links: plans.CourseLinks | None = None) -> plans.CleanData:
    return plans.clean_rows(sources(*rows), links or plans.CourseLinks(set(), {}))


def preview(data: plans.CleanData) -> str:
    return next(iter(plans.render_previews(data["records"]).values()))


def workbook_bytes(headers: tuple[str, ...] = HEADERS, rows: list[plans.RawFields] | None = None) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert isinstance(sheet, Worksheet)
    sheet.append(["培养方案课程列表"])
    sheet.append(headers)
    for fields in rows if rows is not None else [row()]:
        sheet.append([fields.get(header) for header in headers])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class SemesterTests(unittest.TestCase):
    def test_single_and_missing(self):
        self.assertEqual(plans.parse_semesters(1), ["1"])
        self.assertEqual(plans.parse_semesters(" 12 "), ["12"])
        for value in (None, "", "  "):
            self.assertEqual(plans.parse_semesters(value), [])

    def test_ranges_small_terms_and_order(self):
        self.assertEqual(plans.parse_semesters("4,1-3,S1-S2,2,S1"), ["1", "2", "S1", "3", "4", "S2"])
        self.assertEqual(plans.parse_semesters("6,S3"), ["6", "S3"])
        self.assertEqual(plans.parse_semesters("S1,S2,S3"), ["S1", "S2", "S3"])
        self.assertEqual(plans.parse_semesters("10-12,S5"), ["10", "S5", "11", "12"])

    def test_invalid_is_not_silently_discarded(self):
        for value in ("0", "13", "S0", "S6", "4-2", "S3-S1", "S1-3", "1-S2", "1-13", "S1-S6",
                      "1,", ",1", "1,,2", "01", "s1", "1.0", "1-2-3", "1，2", "秋季", "1,bad"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                plans.parse_semesters(value)


class NameTests(unittest.TestCase):
    def test_annotations_and_repeated_suffixes_are_removed(self):
        cases = {
            "人工智能概述(英)": "人工智能概述",
            "传媒调查与统计分析（1）": "传媒调查与统计分析",
            "传媒调查与统计分析（2）": "传媒调查与统计分析",
            "电力系统污秽与覆冰绝缘（研讨课）": "电力系统污秽与覆冰绝缘",
            "计算机辅助翻译（本硕共享）": "计算机辅助翻译",
            "机器学习（本研共享）": "机器学习",
            "中国文化经典（上）（英文）": "中国文化经典",
            "高等数学II-1（全英文）": "高等数学",
            "视听语言Ｉ-2": "视听语言",
            "人体分子与细胞-1st": "人体分子与细胞",
            "人体功能学 - 1": "人体功能学",
            "数智力学项目制课程（二）": "数智力学项目制课程",
            "测试（II－2）（双语课）": "测试",
            "测试（含快题）（拔尖）": "测试",
            "* 测试（（英））": "测试",
            "留学生-初级汉语": "初级汉语",
            "文明经典A": "文明经典系列",
        }
        for original, expected in cases.items():
            with self.subTest(original=original):
                self.assertEqual(plans.simplify_course_name(original), expected)
                self.assertEqual(plans.simplify_course_name(expected), expected)

    def test_core_subjects_and_technical_names_are_not_truncated(self):
        for name in (
            "专业英语", "英语文学", "体育心理学", "机械CAD", "三维CAD", "面向对象技术与UML",
            "材料成型CAD/CAE/CAM", "程序设计基础（基于C++）", "第二外语（日语）",
            "材料专题实验（功能材料）", "建筑类设计理论与方法（建筑学）", "Co-op项目培训",
            "C++程序设计", "3D建模", "图形表达与建筑信息模型（BIM）",
        ):
            with self.subTest(name=name):
                self.assertEqual(plans.simplify_course_name(name), name)

    def test_original_names_render_while_links_use_core_names(self):
        originals = ["新课（1）（英）", "新课（2）（本研共享）"]
        links = plans.CourseLinks(["/course/新课"], {})
        data = clean(*(row(课程名称=name) for name in originals), links=links)
        self.assertEqual([r["name"] for r in data["records"]], originals)
        self.assertEqual({r["raw_fields"]["课程名称"] for r in data["records"]}, set(originals))
        self.assertEqual(data["summary"]["unique_records"], 2)
        for name in originals:
            self.assertIn(f"[{name}](../../../course/新课.md)", preview(data))
        self.assertEqual(clean(row(课程名称="（英）"))["records"][0]["name"], "（英）")
        for name in ("*高等数学II-1（全英文）", "文明经典系列B", "人工智能概述(英)", "计算机辅助翻译（本硕共享）"):
            record = clean(row(课程名称=name))["records"][0]
            self.assertEqual(record["name"], name)
            self.assertIn(plans.markdown_text(name), plans.render_course(record))


class CleaningTests(unittest.TestCase):
    def test_credits(self):
        for value, expected in ((0, "0"), ("0.00", "0"), (2.5, "2.5"), ("2.00", "2"), (None, None), (" ", None)):
            self.assertEqual(plans.normalize_credits(value), expected)
        for value in ("NaN", "Infinity", "-1", "三", "2学分", True):
            with self.subTest(value=value), self.assertRaises(ValueError):
                plans.normalize_credits(value)
        data = clean(row(总学分=0), row(总学分=None))
        self.assertEqual(data["summary"]["zero_credits"], 1)
        self.assertEqual(data["summary"]["missing_credits"], 1)
        self.assertIn(':l-circle-arrow-up:`0`', preview(data))
        self.assertIn(':l-circle-arrow-up:`未提供`', preview(data))

    def test_missing_optional_fields_are_retained(self):
        data = clean(row(课程代码=None, 课程名称=None, 课程类别=None, 课程性质=None, 总学分=None, 开课学期=None))
        output = preview(data)
        self.assertEqual(data["summary"]["unique_records"], 1)
        for label in ("未注明学期", "未注明性质", "未注明类别", "未提供课程名称", ":l-book:`未提供`"):
            self.assertIn(label, output)
        self.assertIsNone(data["records"][0]["total_credits"])
        self.assertEqual(data["records"][0]["semesters"], [])

    def test_missing_identity_fields_fail_with_location(self):
        for field in ("学院", "专业", "年级"):
            with self.subTest(field=field), self.assertRaisesRegex(ValueError, "a.xlsx:3"):
                clean(row(**{field: None}))
        for year in ("NaN", "2026.5", "26", "未知"):
            with self.subTest(year=year), self.assertRaisesRegex(ValueError, "a.xlsx:3"):
                clean(row(年级=year))
        self.assertEqual(clean(row(年级=2026.0))["records"][0]["year"], "2026")

    def test_only_minor_and_dual_degree_are_filtered(self):
        data = clean(
            row(专业="计算机（辅修）", 总学分="无效"), row(专业="双学位计算机"),
            row(层次="研究生"), row(层次="预科"), row(层次=None), row(层次="本科"),
        )
        self.assertEqual(data["summary"]["input_rows"], 6)
        self.assertEqual(data["summary"]["excluded_minor_or_dual_degree"], 2)
        self.assertEqual(data["summary"]["unique_records"], 4)
        self.assertEqual({item["level"] for item in data["records"]}, {"研究生", "预科", "本科", None})

    def test_exact_duplicates_merge_without_resolving_variants(self):
        data = clean(row(), row(), row(备注="仅限本专业"), row(课程性质="选修"),
                     row(课程类别="大类平台课程"), row(总学时=32), row(层次="研究生"))
        self.assertEqual(data["summary"]["duplicate_rows_merged"], 1)
        self.assertEqual(data["summary"]["unique_records"], 6)
        self.assertEqual(sum(len(item["sources"]) for item in data["records"]), 7)
        self.assertEqual(len(next(item for item in data["records"] if item["raw_fields"] == row())["sources"]), 2)
        self.assertTrue(any(item["remarks"] == "仅限本专业" for item in data["records"]))
        self.assertNotIn("备注：", preview(data))
        self.assertIn('=== "选修"', preview(data))

    def test_range_raw_value_and_policy_credits_are_preserved(self):
        data = clean(row(课程名称="形势与政策", 总学分=2, 开课学期="1-8"))
        record = data["records"][0]
        self.assertEqual(record["semester_raw"], "1-8")
        self.assertEqual(record["raw_fields"]["开课学期"], "1-8")
        self.assertEqual(record["total_credits"], "2")
        output = preview(data)
        self.assertEqual(output.count(":l-circle-arrow-up:`总学分 2`"), 8)
        self.assertNotIn("开课范围", output)
        self.assertNotIn("非每学期学分", output)
        self.assertNotIn("0.25", output)
        self.assertIn(":l-circle-arrow-up:`总学分 未提供`", preview(clean(row(总学分=None, 开课学期="S1-S3"))))

    def test_course_collection_displays_without_extra_annotations(self):
        links = plans.CourseLinks(["/course/思政选修课程集"], {})
        data = clean(row(课程名称="思政选修课程集", 课程代码="IPE", 开课学期="2-7", 备注="选1门，修2学分"), links=links)
        record = data["records"][0]
        self.assertEqual(record["remarks"], "选1门，修2学分")
        self.assertEqual(record["semester_raw"], "2-7")
        self.assertEqual(plans.render_course(record), "                * [思政选修课程集](../../../course/思政选修课程集.md) - :l-book:`IPE` - :l-circle-arrow-up:`总学分 2`  ")
        self.assertEqual(preview(data).count("[思政选修课程集]"), 6)

    def test_all_categories_and_unknown_values_are_rendered(self):
        categories = (*plans.CATEGORY_ORDER, "新课程类别")
        data = clean(*(row(课程类别=category) for category in categories), row(课程性质="限选", 课程类别="新课程类别"))
        output = preview(data)
        for category in categories:
            self.assertIn(f"* {category}", output)
        self.assertIn('        === "限选"', output)
        self.assertEqual(output.count("                * "), len(categories) + 1)

    def test_no_old_years_are_merged_and_years_descend(self):
        data = clean(row(年级=2023), row(年级=2026), row(年级=2024), row(年级=2025))
        output = preview(data)
        self.assertEqual(re.findall(r'^=== "(\d+)级"', output, re.M), ["2026", "2025", "2024", "2023"])
        self.assertNotIn("2022级", output)
        self.assertNotIn("2023级", preview(clean(row(年级=2026))))

    def test_deterministic_with_reordered_input(self):
        rows = sources(row(课程代码="A10"), row(课程代码="A2"), row(课程代码="*A2"), row(课程代码="A2"), row(课程代码="A"))
        original = plans.clean_rows(rows, plans.CourseLinks(set(), {}))
        random.Random(42).shuffle(rows)
        shuffled = plans.clean_rows(rows, plans.CourseLinks(set(), {}))
        self.assertEqual(plans.serialize(original), plans.serialize(shuffled))
        self.assertEqual(plans.render_previews(original["records"]), plans.render_previews(shuffled["records"]))
        self.assertLess(preview(original).index(":l-book:`A2`"), preview(original).index(":l-book:`A10`"))

    def test_markdown_escapes_untrusted_text(self):
        data = clean(row(课程名称="[课程]<script>", 课程代码="A`B", 备注="第一行\n<script>*第二行*"))
        output = preview(data)
        self.assertIn(r"\[课程\]\<script\>", output)
        self.assertIn(":l-book:``A`B``", output)
        self.assertNotIn("备注：", output)
        self.assertNotIn("第一行", output)
        self.assertNotIn("<script>", output)
        for major in ("../逃逸", "..", "a\\b", "a\nb"):
            with self.subTest(major=major), self.assertRaises(ValueError):
                clean(row(专业=major))


class LinkTests(unittest.TestCase):
    def setUp(self):
        self.links = plans.CourseLinks(
            {"/course/正确课程", "/course/原始名称", "/course/英语", "/course/高等数学", "/course/有 空格(#)"},
            {"courses": {
                "/course/正确课程": {"codes": ["CODE1", "CONFLICT"]},
                "/course/原始名称": {"codes": ["CONFLICT"]},
                "/course/已删除": {"codes": ["DELETED"]},
            }},
        )

    def test_core_name_precedes_stale_code_and_leading_star_is_supported(self):
        self.assertEqual(self.links.resolve("CODE1", "原始名称"), ("/course/原始名称", "name"))
        self.assertEqual(self.links.resolve("*CODE1", "其他名称"), (None, "unlinked"))
        self.assertEqual(self.links.resolve("*CODE1", None), ("/course/正确课程", "code"))

    def test_names_require_existing_pages_and_do_not_collapse_core_subjects(self):
        self.assertEqual(self.links.resolve(None, "原始名称"), ("/course/原始名称", "name"))
        self.assertEqual(self.links.resolve("NEW", "大学英语III"), (None, "unlinked"))
        self.assertEqual(self.links.resolve("NEW", "高等数学II-1"), ("/course/高等数学", "simplified_name"))
        self.assertEqual(self.links.resolve("DELETED", "已删除"), (None, "unlinked"))
        self.assertEqual(self.links.resolve(None, "不存在III"), (None, "unlinked"))
        self.assertEqual(self.links.resolve("CONFLICT", "原始名称"), ("/course/原始名称", "name"))
        self.assertEqual(self.links.resolve("CONFLICT", "其他名称"), (None, "unlinked"))
        self.assertEqual(self.links.resolve("CONFLICT", None), (None, "ambiguous_code"))
        self.assertNotIn("](", plans.render_course(clean(row(), links=self.links)["records"][0]))

    def test_link_escaping_and_current_icons(self):
        record = clean(row(课程名称="有 空格(#)"), links=self.links)["records"][0]
        output = plans.render_course(record)
        self.assertIn("../../../course/有%20空格%28%23%29.md", output)
        self.assertIn(":l-book:", output)
        self.assertIn(":l-circle-arrow-up:", output)
        self.assertNotIn(":material-", output)


class ExcelTests(unittest.TestCase):
    def read_memory_workbooks(self, files: dict[str, bytes]) -> list[tuple[plans.RawFields, plans.Source]]:
        loader = openpyxl.load_workbook
        with patch.object(Path, "glob", return_value=[Path(name) for name in reversed(files)]), patch.object(
            plans.openpyxl, "load_workbook", side_effect=lambda path, **kwargs: loader(BytesIO(files[path.name]), **kwargs)
        ):
            return list(plans.read_excel_rows(Path("virtual")))

    def test_second_header_row_and_sorted_files(self):
        shuffled_headers = tuple(reversed(HEADERS))
        rows = self.read_memory_workbooks({
            "a.xlsx": workbook_bytes(shuffled_headers, [row(总学分=0, 备注="保留原文")]),
            "z.xlsx": workbook_bytes(),
        })
        self.assertEqual([source["file"] for _, source in rows], ["a.xlsx", "z.xlsx"])
        self.assertEqual(rows[0][0]["总学分"], 0)
        self.assertEqual(rows[0][0]["备注"], "保留原文")
        self.assertEqual(rows[0][1]["row"], 3)

    def test_incorrect_a1_dimensions_do_not_truncate_rows_or_columns(self):
        output = BytesIO()
        with ZipFile(BytesIO(workbook_bytes(rows=[row(备注="完整备注"), row(课程代码="LAST")]))) as source, ZipFile(output, "w") as target:
            for name in source.namelist():
                content = source.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', content)
                target.writestr(name, content)
        rows = self.read_memory_workbooks({"a.xlsx": output.getvalue()})
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0][0]["备注"], "完整备注")
        self.assertEqual(rows[1][0]["课程代码"], "LAST")
        self.assertEqual(rows[1][0]["备注"], None)

    def test_data_beyond_headers_is_not_silently_lost(self):
        content = workbook_bytes()
        workbook = openpyxl.load_workbook(BytesIO(content))
        sheet = workbook.active
        assert isinstance(sheet, Worksheet)
        sheet.cell(3, len(HEADERS) + 1, "不可丢弃")
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        with self.assertRaisesRegex(ValueError, "无表头列含数据"):
            self.read_memory_workbooks({"a.xlsx": output.getvalue()})

    def test_optional_header_missing_is_supported(self):
        headers = tuple(header for header in HEADERS if header not in ("备注", "层次"))
        data = plans.clean_rows(self.read_memory_workbooks({"a.xlsx": workbook_bytes(headers)}), plans.CourseLinks(set(), {}))
        self.assertIsNone(data["records"][0]["level"])
        self.assertIsNone(data["records"][0]["remarks"])

    def test_required_and_duplicate_headers_fail(self):
        for headers in (HEADERS[:-3] + HEADERS[-2:], HEADERS + ("年级",)):
            with self.subTest(headers=headers), self.assertRaisesRegex(ValueError, "a.xlsx:2"):
                plans.read_excel_rows(Path("virtual"))
                self.read_memory_workbooks({"a.xlsx": workbook_bytes(headers)})

    def test_formulas_fail_instead_of_becoming_missing_credits(self):
        with self.assertRaisesRegex(ValueError, "a.xlsx:3 总学分 含公式"):
            self.read_memory_workbooks({"a.xlsx": workbook_bytes(rows=[row(总学分="=1+1")])})

    def test_no_files_fails(self):
        with patch.object(Path, "glob", return_value=[]), self.assertRaisesRegex(ValueError, "未找到 Excel"):
            list(plans.read_excel_rows(Path("virtual")))


class CommandTests(unittest.TestCase):
    def setUp(self):
        data = clean(row())
        self.bundle: plans.Bundle = {**data, "markdown_previews": plans.render_previews(data["records"])}

    def test_default_output_is_only_summary(self):
        output = StringIO()
        with patch.object(plans, "build_bundle", return_value=self.bundle), redirect_stdout(output):
            plans.main([])
        self.assertEqual(json.loads(output.getvalue()), self.bundle["summary"])

    def test_json_and_markdown_are_stdout_only(self):
        for arguments in (["--format", "json"], ["--format", "markdown", "--college", "测试学院", "--major", "测试专业"]):
            output = StringIO()
            with patch.object(plans, "build_bundle", return_value=self.bundle), redirect_stdout(output):
                plans.main(arguments)
            if arguments[1] == "json":
                self.assertEqual(json.loads(output.getvalue()), self.bundle)
            else:
                self.assertEqual(output.getvalue(), preview(self.bundle))

    def test_row_count_mismatch_and_missing_major_fail_without_output(self):
        for arguments in (["--expect-rows", "2"], ["--format", "markdown", "--college", "测试学院", "--major", "旧专业"]):
            output = StringIO()
            with patch.object(plans, "build_bundle", return_value=self.bundle), redirect_stdout(output), redirect_stderr(StringIO()), self.assertRaises(SystemExit):
                plans.main(arguments)
            self.assertEqual(output.getvalue(), "")

    def test_stability_check_detects_changed_output(self):
        changed: plans.Bundle = {**self.bundle, "markdown_previews": {"changed": "changed"}}
        with patch.object(plans, "build_bundle", side_effect=[self.bundle, changed]), redirect_stderr(StringIO()), self.assertRaises(SystemExit):
            plans.main(["--verify-stable"])


class RealDataTests(unittest.TestCase):
    def test_all_rows_and_read_only_deterministic_outputs(self):
        def fingerprint():
            protected = sorted((plans.TOOL_DIR / "data").glob("*.xlsx"))
            protected += sorted(path for path in (plans.REPO_ROOT / "public/doc").rglob("*") if path.is_file())
            protected += [plans.REPO_ROOT / "metadata/course-codes.json"]
            return {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in protected}

        before = fingerprint()
        bundle = plans.build_bundle()
        summary = bundle["summary"]
        self.assertEqual(summary["files"], 4)
        self.assertEqual(summary["input_rows"], 46327)
        self.assertEqual(set(summary["records_by_year"]), {"2023", "2024", "2025", "2026"})
        self.assertEqual(summary["input_rows"], summary["excluded_minor_or_dual_degree"] + summary["duplicate_rows_merged"] + summary["unique_records"])
        self.assertEqual(sum(len(record["sources"]) for record in bundle["records"]), summary["retained_rows"])
        self.assertEqual(set(summary["records_by_category"]), set(plans.CATEGORY_ORDER))
        self.assertGreater(summary["records_by_level"]["研究生"], 0)
        self.assertGreater(summary["zero_credits"], 0)
        self.assertGreater(summary["missing_credits"], 0)
        self.assertGreater(summary["missing_semesters"], 0)
        expected_entries: Counter[str] = Counter()
        for record in bundle["records"]:
            self.assertNotIn("辅修", record["major"])
            self.assertNotIn("双学位", record["major"])
            self.assertEqual(record["name"], plans.text(record["raw_fields"]["课程名称"]))
            self.assertEqual(record["total_credits"], plans.normalize_credits(record["raw_fields"]["总学分"]))
            if record["course_page"]:
                self.assertTrue((plans.REPO_ROOT / "public/doc" / (record["course_page"].lstrip("/") + ".md")).is_file())
            expected_entries[plans.render_course(record)] += max(1, len(record["semesters"]))
        actual_entries = Counter(line for content in bundle["markdown_previews"].values() for line in content.splitlines() if line.startswith("                * "))
        self.assertEqual(expected_entries, actual_entries)
        for relative, content in bundle["markdown_previews"].items():
            for target in re.findall(r"\]\(([^\n]+?)\)", content):
                if target.startswith("../../../course/"):
                    path = plans.REPO_ROOT / "public/doc" / Path(relative).parent / unquote(target)
                    self.assertTrue(path.resolve().is_file(), target)
        digest = hashlib.sha256(plans.serialize(bundle).encode()).hexdigest()
        again = plans.build_bundle()
        self.assertEqual(digest, hashlib.sha256(plans.serialize(again).encode()).hexdigest())
        self.assertEqual(before, fingerprint())


if __name__ == "__main__":
    unittest.main()
